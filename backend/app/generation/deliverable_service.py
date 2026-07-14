"""
Deliverable generation service — uses the ICA LLM client.
Generates POV, executive summary, or roadmap documents from workspace knowledge.
Supports export as Markdown and DOCX.
"""
from __future__ import annotations
import io
import logging

from sqlalchemy.orm import Session

from app.db.models import Concept, ConsultingPattern, Deliverable, Document
from app.retrieval.qa_service import _extract_keywords_llm
from app.graph.builder import build_workspace_graph
from app.graph.traversal import find_nodes_by_keywords, get_neighbourhood
from app.llm import chat
from app.schemas import DeliverableOut, DeliverableResponse, SourceRef

logger = logging.getLogger(__name__)

DELIVERABLE_SYSTEM_PROMPT = """\
You are Bob, an AI-native knowledge compiler for IBM Consulting.
Your role is to draft client-ready consulting deliverables from compiled workspace knowledge.
Using ONLY the knowledge context provided below, generate a structured consulting document.
Do NOT use general knowledge, training data, internet search, or any external information.
Every section must cite its source document inline using [Source: Document Name].
Do NOT introduce facts, claims, or recommendations not supported by the provided context.

The document MUST have these sections (use Markdown headings):
## Executive Summary
## Context & Problem Statement
## IBM Recommendation
## Architecture Considerations
## Roadmap
## Risks & Mitigations

Keep each section concise (3-6 sentences for a POV, 1-2 paragraphs for executive summary, bulleted list for roadmap).
Target audience: {audience}
"""


def _retrieve_context(
    db: Session, workspace_id: int, topic: str | None
) -> tuple[str, list[SourceRef], list[int], list[int]]:
    keywords = _extract_keywords_llm(topic or "key insights and recommendations") if topic else []
    G = build_workspace_graph(db, workspace_id)
    matched_nodes = find_nodes_by_keywords(G, keywords)

    if not matched_nodes:
        matched_nodes = list(G.nodes())

    all_relevant_nodes: set[int] = set()
    for nid in matched_nodes:
        nbr_nodes, _ = get_neighbourhood(G, nid, hops=2)
        all_relevant_nodes.update(nbr_nodes)

    if not all_relevant_nodes and G.nodes():
        all_relevant_nodes = set(G.nodes())

    context_parts: list[str] = []
    source_refs: list[SourceRef] = []
    seen_doc_ids: set[int] = set()
    concept_ids: list[int] = []

    concepts = db.query(Concept).filter(Concept.id.in_(list(all_relevant_nodes))).all()
    for c in concepts:
        concept_ids.append(c.id)
        doc_name = "Unknown"
        if c.source_document_id:
            doc = db.get(Document, c.source_document_id)
            doc_name = doc.title or doc.filename if doc else "Unknown"
        context_parts.append(
            f"[Concept: {c.name}] ({c.type}): {c.description}\nExcerpt: {c.source_excerpt or ''}\nSource: {doc_name}"
        )
        if c.source_document_id and c.source_document_id not in seen_doc_ids:
            seen_doc_ids.add(c.source_document_id)
            source_refs.append(SourceRef(
                document_id=c.source_document_id,
                document_name=doc_name,
                excerpt=c.source_excerpt or "",
            ))

    patterns = db.query(ConsultingPattern).filter(ConsultingPattern.workspace_id == workspace_id).all()
    for p in patterns:
        if any(cid in concept_ids for cid in (p.related_concept_ids or [])):
            steps = "\n".join(f"  - {s}" for s in (p.ibm_approach or []))
            context_parts.append(f"[Pattern: {p.name}]\nProblem: {p.problem_statement}\nApproach:\n{steps}")
            for did in (p.source_document_ids or []):
                if did not in seen_doc_ids:
                    doc = db.get(Document, int(did))
                    if doc:
                        seen_doc_ids.add(did)
                        source_refs.append(SourceRef(
                            document_id=did,
                            document_name=doc.title or doc.filename,
                            excerpt="(Pattern context)",
                        ))

    return "\n\n---\n\n".join(context_parts), source_refs, concept_ids, list(seen_doc_ids)


def generate_deliverable(
    db: Session,
    workspace_id: int,
    deliverable_type: str,
    topic: str | None,
    audience: str | None,
) -> DeliverableResponse:
    context_text, source_refs, concept_ids, doc_ids = _retrieve_context(db, workspace_id, topic)

    if not context_text.strip():
        raise ValueError("No knowledge found in workspace to generate a deliverable. Upload and process documents first.")

    topic_str = topic or "the workspace knowledge"
    audience_str = audience or "consulting executive"
    type_label = {
        "POV": "Point of View",
        "executive_summary": "Executive Summary",
        "roadmap": "Strategic Roadmap",
    }.get(deliverable_type, deliverable_type)
    title = f"{type_label}: {topic_str}"

    system = DELIVERABLE_SYSTEM_PROMPT.format(audience=audience_str)
    user_prompt = (
        f"Generate a {type_label} document about: {topic_str}\n\n"
        f"Knowledge Context:\n{context_text}"
    )

    try:
        content_markdown = chat(
            system=system,
            user=user_prompt,
            max_tokens=4096,
        )
    except Exception as e:
        logger.error(f"Deliverable generation failed: {e}")
        raise RuntimeError(f"Deliverable generation failed: {e}")

    deliverable = Deliverable(
        workspace_id=workspace_id,
        type=deliverable_type,
        title=title,
        content_markdown=content_markdown,
        source_concept_ids=concept_ids,
        source_document_ids=doc_ids,
    )
    db.add(deliverable)
    db.commit()
    db.refresh(deliverable)

    return DeliverableResponse(
        deliverable=DeliverableOut.model_validate(deliverable),
        sources=source_refs,
    )


def export_as_docx(content_markdown: str, title: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    title_para = doc.add_heading(title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for line in content_markdown.split("\n"):
        if line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
