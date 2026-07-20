"""
Tests for the geometric / text-fit visual-QA analyzer (tools/visual_qa.py).

These prove the detectors actually fire — a QA tool that never flags anything is
worthless — using python-pptx primitives directly (no IPC template needed).
"""
from __future__ import annotations

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from tools.visual_qa import analyze_prs


def _blank_slide():
    prs = Presentation()  # default 10" x 7.5" canvas
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    return prs, slide


def test_offslide_shape_flagged_as_error():
    prs, slide = _blank_slide()
    # Place a text box entirely past the right edge of the slide.
    tb = slide.shapes.add_textbox(Emu(prs.slide_width + Inches(1)), Emu(0), Inches(3), Inches(1))
    tb.text_frame.text = "Off the edge"
    r = analyze_prs(prs)
    assert any(kind == "off_slide" and sev == "ERROR" for _, sev, kind, _ in r["findings"])
    assert r["errors"]


def test_text_overflow_estimate_flagged():
    prs, slide = _blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(0.4))  # tiny box
    tb.text_frame.text = ("This is a very long paragraph that cannot possibly fit inside a "
                          "one inch by four tenths of an inch text box at a normal font size, "
                          "so the estimator must flag it as overflowing its own frame. ") * 2
    r = analyze_prs(prs)
    assert any(kind == "text_overflow_estimate" for _, _, kind, _ in r["findings"])


def test_clean_shape_produces_no_findings():
    prs, slide = _blank_slide()
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run(); run.text = "A short, well-fitted headline."
    run.font.size = Pt(18)
    r = analyze_prs(prs)
    assert r["findings"] == []
