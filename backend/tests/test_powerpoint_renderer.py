"""
Renderer coverage — exercises the REAL PowerPoint generator (no mock).

This is the automated regression net for the platform's last line of defense
against template text leaking into client decks: generate_pptx, _scan_leaks,
_scrub_template_text, and _looks_like_leak. Prior to this file the entire
renderer was mocked in every test.

Template-dependent tests skip cleanly when IPC_PPT_Template_2026.pptx is absent
(it is a large binary that may not be present in every checkout / CI runner).
"""
from __future__ import annotations

import io

import pytest

import deliverables.generators.powerpoint_generator as ppg
from deliverables.generators.powerpoint_generator import (
    IPC_TEMPLATE_PATH,
    _iter_all_shapes,
    _looks_like_leak,
    _scan_empty_placeholders,
    _scan_leaks,
    generate_pptx,
)

pytestmark = pytest.mark.filterwarnings("ignore::UserWarning")

_TEMPLATE_AVAILABLE = IPC_TEMPLATE_PATH.exists()
_needs_template = pytest.mark.skipif(
    not _TEMPLATE_AVAILABLE, reason="IPC template not present in this checkout"
)


# ── Unit: leak-signature detector (no template needed) ─────────────────────────

@pytest.mark.parametrize("text", [
    "Lorem ipsum dolor sit amet",
    "Click to edit Master text styles",
    "68/84/132pt title",
    "28/36/42pt headline",
    "168pt Go simple and big",
    "Firstname Lastname",
    "email.address@ibm.com",
    "Source: If applicable, describe source origin",
    "Source: 1. Lorem source name",
    "Numbers are optional. Lorem ipsum",
    "Section one",
])
def test_looks_like_leak_flags_template_signatures(text):
    assert _looks_like_leak(text) is True


@pytest.mark.parametrize("text", [
    "Stablecoins will reach $1.9 trillion by 2030.",
    "IBM Digital Asset Haven provides custody across 42+ blockchains.",
    "Institutions should establish custody governance before scaling tokenization.",
    "",
    "   ",
])
def test_looks_like_leak_passes_real_content(text):
    assert _looks_like_leak(text) is False


# ── Helpers ────────────────────────────────────────────────────────────────────

def _all_text(prs) -> str:
    out = []
    for s in prs.slides:
        for sh in _iter_all_shapes(s.shapes):
            try:
                if sh.has_text_frame:
                    out.append(sh.text_frame.text)
            except Exception:
                pass
    return "\n".join(out).lower()


def _sources_slide_count(prs) -> int:
    n = 0
    for s in prs.slides:
        for sh in s.shapes:
            try:
                if sh.has_text_frame and sh.text_frame.text.strip().lower().startswith("sources"):
                    n += 1
                    break
            except Exception:
                pass
    return n


def _full_deck() -> dict:
    return {
        "deliverable_type": "client_201",
        "title": "Digital assets are reshaping settlement economics for global banks",
        "metadata": {"source_documents": ["Doc A", "Doc B", "Doc A"]},
        "slides": [
            {"title": "Blockchain settlement eliminates the capital trap", "layout": "large_text"},
            {"title": "Three mandates are accelerating adoption", "layout": "section_divider"},
            {"title": "Custody, treasury, and settlement form the core", "layout": "title_content",
             "bullets": ["Custody secures assets end to end.", "Treasury automates liquidity."]},
            {"title": "Current versus target state diverge sharply", "layout": "two_col_dividers",
             "col_heads": ["Today", "Target"],
             "columns": [["Batch settlement runs overnight."], ["Real-time settlement clears instantly."]]},
            {"title": "IBM captures value across six patterns", "layout": "six_boxes",
             "boxes": ["Custody.", "Tokenization.", "Settlement.", "Treasury.", "Stablecoins.", "Compliance."]},
            {"title": "Four capability gaps constrain institutions", "layout": "four_boxes_wide",
             "boxes": ["Custody immature.", "Interop limited.", "Governance unclear.", "Talent scarce."]},
            {"title": "Three pillars anchor the operating model", "layout": "four_column_headlines",
             "col_heads": ["Secure", "Scale", "Govern"],
             "columns": [["Key custody protects assets."], ["Throughput scales."], ["Policy enforces compliance."]]},
            {"title": "Adoption climbs while cost falls", "layout": "data_2_callouts",
             "stats": [{"label": "+240%", "body": "Volume growth."}, {"label": "-70%", "body": "Cost reduction."}]},
            {"title": "Modernization proceeds through five stages", "layout": "process_diagram",
             "steps": [{"heading": "Assess", "body": "Baseline custody."},
                       {"heading": "Design", "body": "Define architecture."},
                       {"heading": "Pilot", "body": "Prove one corridor."}]},
            # diagram layouts whose template carries baked-in demo text (scrub must clean)
            {"title": "The reference architecture spans custody and settlement",
             "layout": "technical_architecture"},
            {"title": "The roadmap delivers value in three horizons", "layout": "timeline"},
            {"title": "Capabilities decompose into a clear hierarchy", "layout": "hierarchy",
             "bullets": ["Custody underpins everything.", "Compliance spans all layers."]},
            {"title": "Value decomposes from speed to freed capital", "layout": "value_tree",
             "bullets": ["Faster settlement frees capital.", "Freed capital funds lending."]},
            {"title": "Governance assigns clear accountability", "layout": "raci",
             "bullets": ["The CISO owns custody security.", "Compliance owns reporting."]},
            {"title": "Six to ten items split across two columns", "layout": "two_column",
             "bullets": ["Point one.", "Point two.", "Point three.", "Point four."]},
            # an LLM-emitted stray sources slide — the renderer must not double it
            {"title": "Sources", "layout": "title_content", "bullets": ["Should be skipped."]},
        ],
    }


# ── generate_pptx end-to-end ───────────────────────────────────────────────────

@_needs_template
def test_generate_pptx_all_layouts_no_leaks():
    from pptx import Presentation
    data = generate_pptx(_full_deck(), "Test Workspace")
    assert isinstance(data, (bytes, bytearray)) and len(data) > 10_000
    prs = Presentation(io.BytesIO(data))
    # The scan is also run internally (hard-fail); this re-confirms on reload.
    assert _scan_leaks(prs) == []
    assert len(prs.slides) > 0


@_needs_template
def test_generate_pptx_single_sources_slide():
    from pptx import Presentation
    prs = Presentation(io.BytesIO(generate_pptx(_full_deck(), "Test Workspace")))
    assert _sources_slide_count(prs) == 1  # LLM 'Sources' skipped; one appended


@_needs_template
def test_cover_title_rendered_not_leaked():
    from pptx import Presentation
    deck = _full_deck()
    prs = Presentation(io.BytesIO(generate_pptx(deck, "Test Workspace")))
    txt = _all_text(prs)
    assert "68/84" not in txt and "firstname lastname" not in txt
    assert "digital assets are reshaping settlement economics" in txt  # real cover title


@_needs_template
def test_four_column_headlines_title_not_leaked():
    from pptx import Presentation
    deck = {"deliverable_type": "client_101", "metadata": {"source_documents": []},
            "slides": [{"title": "Three pillars anchor the operating model",
                        "layout": "four_column_headlines",
                        "col_heads": ["A", "B", "C"],
                        "columns": [["One."], ["Two."], ["Three."]]}]}
    prs = Presentation(io.BytesIO(generate_pptx(deck, "WS")))
    txt = _all_text(prs)
    assert "28/36/42pt headline" not in txt
    assert "three pillars anchor the operating model" in txt


@_needs_template
def test_process_diagram_unused_boxes_blanked():
    from pptx import Presentation
    deck = {"deliverable_type": "client_201", "metadata": {"source_documents": []},
            "slides": [{"title": "Five-stage modernization", "layout": "process_diagram",
                        "steps": [{"heading": "Assess", "body": "Baseline."},
                                  {"heading": "Design", "body": "Architect."},
                                  {"heading": "Pilot", "body": "Prove."}]}]}
    prs = Presentation(io.BytesIO(generate_pptx(deck, "WS")))
    assert "lorem" not in _all_text(prs)  # 4 unused step boxes must be blanked


@_needs_template
def test_no_empty_placeholders_survive():
    """No placeholder renders 'Click to add text' — every empty placeholder is
    removed, not blanked."""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(generate_pptx(_full_deck(), "Test Workspace")))
    assert _scan_empty_placeholders(prs) == []


@_needs_template
def test_deck_terminates_on_sources_no_blank_end_slide():
    """The deck ends on the Sources slide — no blank/decorative terminal slide."""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(generate_pptx(_full_deck(), "Test Workspace")))
    last = list(prs.slides)[-1]
    last_text = " ".join(
        sh.text_frame.text for sh in _iter_all_shapes(last.shapes)
        if getattr(sh, "has_text_frame", False)
    ).strip().lower()
    assert last_text.startswith("sources")


@_needs_template
def test_cover_has_no_empty_contact_placeholders():
    """The cover's unused contact placeholders are removed (not left empty)."""
    from pptx import Presentation
    prs = Presentation(io.BytesIO(generate_pptx(_full_deck(), "Test Workspace")))
    cover = list(prs.slides)[0]
    empties = [sh.name for sh in cover.shapes
               if getattr(sh, "is_placeholder", False) and sh.has_text_frame
               and sh.text_frame.text.strip() == ""]
    assert empties == []


@_needs_template
def test_generate_pptx_hard_fails_when_placeholder_removal_disabled(monkeypatch):
    """If empty-placeholder removal is broken, the post-render QA gate must ABORT
    rather than ship a deck full of 'Click to add text' prompts."""
    monkeypatch.setattr(ppg, "_remove_empty_placeholders", lambda slide: 0)
    with pytest.raises(ValueError, match="empty placeholder"):
        generate_pptx(_full_deck(), "Test Workspace")


@_needs_template
def test_generate_pptx_hard_fails_when_scrub_disabled(monkeypatch):
    """If scrubbing is broken, a diagram layout's baked-in 'Lorem ipsum' survives
    and the internal leak scan must ABORT generation rather than ship it."""
    monkeypatch.setattr(ppg, "_scrub_template_text", lambda slide: 0)
    deck = {"deliverable_type": "client_201", "metadata": {"source_documents": []},
            "slides": [{"title": "Capabilities decompose into a hierarchy", "layout": "hierarchy",
                        "bullets": ["Custody underpins everything."]}]}
    with pytest.raises(ValueError, match="placeholder"):
        generate_pptx(deck, "WS")
